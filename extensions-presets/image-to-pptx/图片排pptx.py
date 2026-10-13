import os
import sys
import subprocess
from datetime import datetime
from pathlib import Path

# ================= 自动安装依赖库 =================
try:
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image, ImageOps
except ModuleNotFoundError:
    print("检测到未安装 pptx 库，正在自动安装，请稍候...")
    try:
        # 使用国内清华镜像源加速下载，防止网络卡顿
        subprocess.check_call([
            sys.executable, "-m", "pip", "install", "python-pptx",
            "-i", "https://pypi.tuna.tsinghua.edu.cn/simple"
        ])
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image, ImageOps
        print("pptx 库安装成功！继续执行脚本...\n")
    except Exception as e:
        print(f"自动安装失败，请检查网络连接。错误信息: {e}")
        sys.exit(1)
# =================================================

import glob

IMG_EXTENSIONS = ("*.jpg", "*.jpeg", "*.png", "*.JPG", "*.JPEG", "*.PNG")


def create_pptx_from_images():
    # 输入/输出一律相对当前工作目录：手动双击时即脚本所在文件夹；
    # 被 CanvasFlow 扩展服务调用时，工作区里有 input/，输出进 output/（见 拓展脚本规范.md）
    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    image_files_set = set()
    for ext in IMG_EXTENSIONS:
        for file in glob.glob(os.path.join("input", ext)):
            image_files_set.add(os.path.abspath(file))
    image_files = sorted(image_files_set)

    if not image_files:
        print("input 文件夹下没有找到图片文件（支持 jpg/png）")
        sys.exit(1)

    print(f"共发现 {len(image_files)} 张图片，开始导入...")

    blank_layout = prs.slide_layouts[6]
    for img_path in image_files:
        slide = prs.slides.add_slide(blank_layout)
        # 等比缩放：高度一律填满页高，不拉伸变形，宽度按比例居中（不足或超出均居中）
        with Image.open(img_path) as im:
            img_w, img_h = ImageOps.exif_transpose(im).size
        scale = prs.slide_height / img_h
        pic_w = round(img_w * scale)
        pic_h = round(img_h * scale)
        slide.shapes.add_picture(
            img_path,
            left=round((prs.slide_width - pic_w) / 2),
            top=round((prs.slide_height - pic_h) / 2),
            width=pic_w,
            height=pic_h
        )
        print(f"已置入: {os.path.basename(img_path)}")

    output_name = f"图片置入pptx_{datetime.now().strftime('%Y%m%d-%H%M%S')}.pptx"
    prs.save(output_dir / output_name)
    print(f"成功！PPTX 已生成: {(output_dir / output_name).resolve()}")


if __name__ == "__main__":
    create_pptx_from_images()
